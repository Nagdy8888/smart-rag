# pyright: reportUnknownMemberType=false, reportUnknownVariableType=false, reportAttributeAccessIssue=false
"""PPTX adapter with versioned reading-order decisions."""

from __future__ import annotations

import time
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from smart_rag_core.ingestion.parser.exceptions import ParserLimitExceeded
from smart_rag_core.ingestion.parser.helpers import (
    as_float,
    as_int,
    build_coverage,
    element_id,
    vendor_attr,
)
from smart_rag_core.ingestion.parser.limits import (
    CLEANING_POLICY_VERSION,
    PARSER_CONFIG_VERSION,
    PPTX_MAX_SLIDES,
)
from smart_rag_core.ingestion.parser.types import (
    ElementKind,
    ParseDisposition,
    ParseResult,
    ParserResourceMetrics,
    PptxLocator,
    TypedElement,
)


def parse_pptx(path: Path, content_hash: str, format_key: str) -> ParseResult:
    """Parse slides/shapes with reading-order evidence separate from z-order."""

    started = time.monotonic()
    presentation = Presentation(str(path))
    if len(presentation.slides) > PPTX_MAX_SLIDES:
        raise ParserLimitExceeded("slides", f"slide count {len(presentation.slides)}")
    elements: list[TypedElement] = []
    order = 0
    for slide_index, slide in enumerate(presentation.slides):
        elements.append(
            TypedElement(
                element_id=element_id(
                    content_hash,
                    order,
                    ElementKind.SLIDE_BREAK,
                    PptxLocator(
                        slide_index=slide_index,
                        reading_order_index=order,
                    ),
                    None,
                ),
                kind=ElementKind.SLIDE_BREAK,
                text=None,
                reading_order=order,
                locator=PptxLocator(slide_index=slide_index, reading_order_index=order),
            )
        )
        order += 1
        shape_entries: list[tuple[float, float, int, object]] = []
        for shape in slide.shapes:
            top = as_float(vendor_attr(shape, "top", 0) or 0)
            left = as_float(vendor_attr(shape, "left", 0) or 0)
            z_order = as_int(vendor_attr(shape, "shape_id", 0) or 0)
            shape_entries.append((top, left, z_order, shape))
        for reading_index, (_top, _left, z_order, shape) in enumerate(
            sorted(shape_entries, key=lambda item: (item[0], item[1], item[2]))
        ):
            if vendor_attr(shape, "shape_type") == MSO_SHAPE_TYPE.GROUP:
                elements.append(
                    TypedElement(
                        element_id=element_id(
                            content_hash,
                            order,
                            ElementKind.FIGURE,
                            PptxLocator(
                                slide_index=slide_index,
                                shape_id=_shape_id(shape),
                                z_order=z_order,
                                reading_order_index=reading_index,
                            ),
                            None,
                        ),
                        kind=ElementKind.FIGURE,
                        text=None,
                        reading_order=order,
                        locator=PptxLocator(
                            slide_index=slide_index,
                            shape_id=_shape_id(shape),
                            z_order=z_order,
                            reading_order_index=reading_index,
                            geometry_x=_shape_emu(shape, "left"),
                            geometry_y=_shape_emu(shape, "top"),
                            geometry_w=_shape_emu(shape, "width"),
                            geometry_h=_shape_emu(shape, "height"),
                        ),
                        media_only=True,
                        metadata={
                            "unsupported_region": True,
                            "smartart_or_group": True,
                        },
                    )
                )
                order += 1
                continue
            text = ""
            if hasattr(shape, "text"):
                text = str(vendor_attr(shape, "text")).strip()
            if not text:
                if vendor_attr(shape, "shape_type") in {
                    MSO_SHAPE_TYPE.PICTURE,
                    MSO_SHAPE_TYPE.CHART,
                }:
                    elements.append(
                        TypedElement(
                            element_id=element_id(
                                content_hash,
                                order,
                                ElementKind.FIGURE,
                                PptxLocator(
                                    slide_index=slide_index,
                                    shape_id=_shape_id(shape),
                                    reading_order_index=reading_index,
                                ),
                                None,
                            ),
                            kind=ElementKind.FIGURE,
                            text=None,
                            reading_order=order,
                            locator=PptxLocator(
                                slide_index=slide_index,
                                shape_id=_shape_id(shape),
                                z_order=z_order,
                                reading_order_index=reading_index,
                            ),
                            media_only=True,
                        )
                    )
                    order += 1
                continue
            locator = PptxLocator(
                slide_index=slide_index,
                shape_id=_shape_id(shape),
                z_order=z_order,
                reading_order_index=reading_index,
                geometry_x=_shape_emu(shape, "left"),
                geometry_y=_shape_emu(shape, "top"),
                geometry_w=_shape_emu(shape, "width"),
                geometry_h=_shape_emu(shape, "height"),
            )
            elements.append(
                TypedElement(
                    element_id=element_id(
                        content_hash, order, ElementKind.PARAGRAPH, locator, text
                    ),
                    kind=ElementKind.PARAGRAPH,
                    text=text,
                    reading_order=order,
                    locator=locator,
                )
            )
            order += 1

    elements_tuple = tuple(elements)
    config_hash = ParseResult.stable_config_hash(
        PARSER_CONFIG_VERSION, CLEANING_POLICY_VERSION
    )
    return ParseResult(
        disposition=ParseDisposition.SUPPORTED,
        format=format_key,
        content_hash=content_hash,
        parser_version=PARSER_CONFIG_VERSION,
        cleaning_version=CLEANING_POLICY_VERSION,
        config_hash=config_hash,
        elements=elements_tuple,
        coverage=build_coverage(elements_tuple),
        resource_metrics=ParserResourceMetrics(
            wall_seconds=time.monotonic() - started,
            pages_processed=len(presentation.slides),
        ),
        detail="pptx parsed",
    )


def _shape_id(shape: object) -> int:
    """Return ``int(shape.shape_id)`` for an untyped PPTX shape.

    Args:
        shape: python-pptx shape handle stored as ``object``.

    Returns:
        Integer shape id.

    Raises:
        TypeError: If ``shape_id`` is not convertible to int.
        ValueError: If ``shape_id`` is a non-integer string.
    """

    return as_int(vendor_attr(shape, "shape_id"))


def _shape_emu(shape: object, name: str) -> float:
    """Return ``float(shape.<name>)`` for an untyped PPTX geometry field.

    Args:
        shape: python-pptx shape handle stored as ``object``.
        name: Attribute name such as ``left``, ``top``, ``width``, or ``height``.

    Returns:
        Geometry value in EMUs as float.

    Raises:
        TypeError: If the attribute is not convertible to float.
        ValueError: If the attribute is a non-numeric string.
    """

    return as_float(vendor_attr(shape, name))
