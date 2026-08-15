# pyright: reportUnknownMemberType=false, reportUnknownVariableType=false, reportAttributeAccessIssue=false
"""XLSX adapter preserving typed cells, formulas, and cache authority."""

from __future__ import annotations

import time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

from smart_rag_core.ingestion.parser.exceptions import (
    ParserEncryptedError,
    ParserLimitExceeded,
)
from smart_rag_core.ingestion.parser.helpers import build_coverage, element_id
from smart_rag_core.ingestion.parser.limits import (
    CLEANING_POLICY_VERSION,
    PARSER_CONFIG_VERSION,
    PARSER_SANDBOX_MAX_ROWS,
    PARSER_SANDBOX_MAX_SHEETS,
)
from smart_rag_core.ingestion.parser.types import (
    ElementKind,
    ParseDisposition,
    ParseResult,
    ParserResourceMetrics,
    TypedElement,
    XlsxLocator,
)


def parse_xlsx(path: Path, content_hash: str, format_key: str) -> ParseResult:
    """Parse workbook sheets without recalculating formulas."""

    started = time.monotonic()
    try:
        workbook = load_workbook(path, data_only=False)
    except Exception as exc:
        if "password" in str(exc).lower() or "encrypted" in str(exc).lower():
            raise ParserEncryptedError("encrypted workbook") from exc
        raise
    if len(workbook.sheetnames) > PARSER_SANDBOX_MAX_SHEETS:
        raise ParserLimitExceeded("sheets", f"sheet count {len(workbook.sheetnames)}")
    epoch: str = "1904" if workbook.excel_base_date.year == 1904 else "1900"
    elements: list[TypedElement] = []
    order = 0
    for sheet_index, sheet_name in enumerate(workbook.sheetnames):
        sheet = workbook[sheet_name]
        hidden = sheet.sheet_state in {"hidden", "veryHidden"}
        row_count = 0
        for row in sheet.iter_rows():
            row_count += 1
            if row_count > PARSER_SANDBOX_MAX_ROWS:
                raise ParserLimitExceeded(
                    "rows", f"row count exceeds {PARSER_SANDBOX_MAX_ROWS}"
                )
            for cell in row:
                if not isinstance(cell, Cell):
                    continue
                value_text, metadata = _cell_fields(cell)
                if value_text is None and not metadata:
                    continue
                locator = XlsxLocator(
                    sheet_name=sheet_name,
                    sheet_index=sheet_index,
                    row=int(cell.row),
                    col=int(cell.column),
                    cell_ref=str(cell.coordinate),
                    hidden_sheet=hidden,
                    date_epoch=epoch,  # type: ignore[arg-type]
                    number_format=str(cell.number_format),
                )
                elements.append(
                    TypedElement(
                        element_id=element_id(
                            content_hash,
                            order,
                            ElementKind.TABLE_CELL,
                            locator,
                            value_text,
                        ),
                        kind=ElementKind.TABLE_CELL,
                        text=value_text,
                        reading_order=order,
                        locator=locator,
                        metadata=metadata,
                    )
                )
                order += 1
    workbook.close()
    elements_tuple = tuple(elements)
    config_hash = ParseResult.stable_config_hash(
        PARSER_CONFIG_VERSION, CLEANING_POLICY_VERSION
    )
    return ParseResult(
        disposition=ParseDisposition.SUPPORTED,
        format=format_key,
        content_hash=content_hash,
        parser_version=PARSER_CONFIG_VERSION,
        cleaning_version=CLEANING_POLICY_VERSION,
        config_hash=config_hash,
        elements=elements_tuple,
        coverage=build_coverage(elements_tuple),
        resource_metrics=ParserResourceMetrics(wall_seconds=time.monotonic() - started),
        detail="xlsx parsed",
    )


def _cell_fields(
    cell: Cell,
) -> tuple[str | None, dict[str, str | int | float | bool | None]]:
    """Extract typed value, formula text, and cache authority metadata."""

    metadata: dict[str, str | int | float | bool | None] = {}
    formula = getattr(cell, "value", None)
    if getattr(cell, "data_type", None) == "f" or (
        isinstance(formula, str) and formula.startswith("=")
    ):
        metadata["formula_text"] = str(formula)
        metadata["cache_authority"] = "unevaluated"
        return str(formula), metadata
    cached = cell.value
    if cached is None:
        return None, metadata
    metadata["cache_authority"] = "stored"
    if isinstance(cached, bool):
        metadata["typed_value_kind"] = "boolean"
        return str(cached).lower(), metadata
    if isinstance(cached, int):
        metadata["typed_value_kind"] = "integer"
        return str(cached), metadata
    if isinstance(cached, float):
        metadata["typed_value_kind"] = "float"
        return repr(cached), metadata
    metadata["typed_value_kind"] = "string"
    return str(cached), metadata
